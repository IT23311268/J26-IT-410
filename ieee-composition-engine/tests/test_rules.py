"""
Tests for the IEEE rulebook (src/rules/ieee_rules.py) and the rules added with it.

Run with:   pytest -q

They check three things:
  1. the rulebook itself is complete and consistent (unique IDs, real
     sources, real files, every score check tied to a rule)
  2. the engine really reads its values from the rulebook
  3. every new rule does what the rulebook says
"""

from __future__ import annotations

import os
import sys

import docx
import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from src.compose.front_matter import apply_front_matter_rules, clean_abstract   # noqa: E402
from src.compose.ieee_text import apply_text_rules, number_style_counts          # noqa: E402
from src.compose.references import abbreviate_venue, format_ieee, segments_to_text  # noqa: E402
from src.evaluate import compliance                                              # noqa: E402
from src.models import Element, Section, StructuredPaper                         # noqa: E402
from src.render import ieee_style as S                                           # noqa: E402
from src.render import render_docx                                               # noqa: E402
from src.rules import ieee_rules as R                                            # noqa: E402
from src.structure import map_to_ieee                                            # noqa: E402


# ---------------------------------------------------------------------------
# 1. The rulebook is complete and consistent
# ---------------------------------------------------------------------------

def test_every_rule_is_well_formed():
    ids = [r.id for r in R.RULES]
    assert len(ids) == len(set(ids)), "rule IDs must be unique"
    for r in R.RULES:
        assert r.source in R.SOURCES, f"{r.id}: unknown source {r.source}"
        assert r.status in (R.ENFORCED, R.CHECKED, R.AUTHOR), r.id
        assert r.rule.strip(), r.id
        if r.status in (R.ENFORCED, R.CHECKED):
            assert r.where, f"{r.id}: an applied or checked rule must say which file does it"
        for path in filter(None, r.where.split(",")):
            assert os.path.exists(os.path.join(ROOT, path.strip())), f"{r.id}: {path} missing"


def test_score_checks_and_rules_point_at_each_other():
    keys = {key for key, _, _ in compliance.checks()}
    for r in R.RULES:
        if r.check:
            assert r.check in keys, f"{r.id} names a check that does not exist: {r.check}"
    for key in keys:
        assert R.rules_for_check(key), f"score check '{key}' does not test any rulebook rule"


def test_the_readable_rulebook_is_up_to_date():
    sys.path.insert(0, os.path.join(ROOT, "tools"))
    import export_rules
    with open(export_rules.OUT, encoding="utf-8") as f:
        assert f.read() == export_rules.build(), "run: python tools/export_rules.py"


# ---------------------------------------------------------------------------
# 2. The engine reads its values from the rulebook
# ---------------------------------------------------------------------------

def test_layout_values_come_from_the_rulebook():
    S.use_paper("letter")
    assert S.TITLE_SIZE.pt == R.SIZE_PT["title"] == 24
    assert S.BODY_SIZE.pt == R.SIZE_PT["body"] == 10
    assert S.MARGIN_TOP.inches == R.PAPER_SIZES["letter"].margin_top_in == 0.75
    assert S.FONT == R.FONT and S.ABSTRACT_LABEL == R.LABELS["abstract"]


def test_a4_paper_size(tmp_path):
    try:
        S.use_paper("a4")
        assert round(S.PAGE_WIDTH.mm) == 210 and round(S.PAGE_HEIGHT.mm) == 297
        path = str(tmp_path / "a4.docx")
        render_docx(StructuredPaper(source_path="t", title="A4 Paper", abstract="Text."), path)
        section = docx.Document(path).sections[0]
        assert round(section.page_width.mm) == 210
        assert round(section.top_margin.mm) == 19
    finally:
        S.use_paper("letter")
    assert S.PAGE_WIDTH.inches == 8.5
    with pytest.raises(ValueError):
        S.use_paper("legal")


# ---------------------------------------------------------------------------
# 3. The new rules
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("before, after", [
    ("an accuracy of .95 was reached", "an accuracy of 0.95 was reached"),        # NUM-01
    ("a 10mm nodule and a 2.5GHz clock", "a 10 mm nodule and a 2.5 GHz clock"),  # NUM-02
    ("improved by 4 %", "improved by 4%"),                                        # NUM-03
    ("between 20-30% of cases", "between 20%–30% of cases"),                      # NUM-04
    ("nodules of 3-30 mm", "nodules of 3–30 mm"),                                 # NUM-04
    ("the H100 GPU and version 2.1 of GPT-4", "the H100 GPU and version 2.1 of GPT-4"),  # untouched
])
def test_number_and_unit_rules(before, after):
    assert apply_text_rules(before)[0] == after


def test_number_style_counts():
    assert number_style_counts("0.95 and 10 mm and 4%") == (3, 0)
    assert number_style_counts(".95 and 10mm") == (0, 2)


def test_figure_rules():
    assert apply_text_rules("as Figure 3a shows")[0] == "as Fig. 3(a) shows"             # FIG-04
    assert apply_text_rules("see Figure 3 (b).")[0] == "see Fig. 3(b)."
    assert apply_text_rules("Figures 2 and 3 compare")[0] == "Figs. 2 and 3 compare"
    assert apply_text_rules("Tables 2 and 3 list")[0] == "Tables II and III list"


def test_section_cross_references():
    numbers = {"3": "III", "3.1": "III-A", "4": "IV"}
    text, counts = apply_text_rules("as in Section 3.1 and Subsection 3.1, see Sec. 4.", numbers)
    assert text == "as in Section III-A and Section III-A, see Section IV."             # HDG-09
    assert counts["section_ref"] == 3
    assert apply_text_rules("in Sections 3 and 4", numbers)[0] == "in Sections III and IV"


def test_citation_rules():
    assert apply_text_rules("as shown in reference [3].")[0] == "as shown in [3]."      # CIT-05
    assert apply_text_rules("This was shown. [3] The next")[0] == \
        "This was shown [3]. The next"                                                  # CIT-04
    assert apply_text_rules("Setio et al. [3] proposed")[0] == "Setio et al. [3] proposed"


def test_contractions_and_spacing():
    assert apply_text_rules("We don't know.  It's new.")[0] == "We do not know. It is new."


@pytest.mark.parametrize("venue, short", [
    ("IEEE Transactions on Medical Imaging", "IEEE Trans. Med. Imag."),
    ("Journal of Machine Learning Research", "J. Mach. Learn. Res."),
    ("Advances in Neural Information Processing Systems", "Adv. Neural Inf. Process. Syst."),
    ("Nature", "Nature"),                                     # one word: never abbreviated
])
def test_venue_abbreviations(venue, short):
    assert abbreviate_venue(venue) == short                   # REF-06


def test_reference_without_a_year_gets_nd():
    record = {"authors": [{"given": "Ann", "family": "Lee"}], "title": "A tool",
              "venue": "Software Report", "year": ""}
    assert R.NO_DATE in segments_to_text(format_ieee(record))  # REF-10


def test_front_matter_rules():
    paper = StructuredPaper(source_path="t", title="DEEP LEARNING FOR CT SCANS (Invited)",
                            authors=["J.K. Smith", "A. Lee"],
                            abstract="Prior detectors [3], [4] miss small nodules "
                                     "(Smith et al., 2020). Results are strong.")
    apply_front_matter_rules(paper)
    assert paper.title == "Deep Learning for CT Scans"                  # TTL-01, TTL-02
    assert paper.authors == ["J. K. Smith", "A. Lee"]                   # TTL-03
    assert paper.abstract == "Prior detectors miss small nodules. Results are strong."  # ABS-04
    assert clean_abstract("Shown by Smith et al. (2019) earlier.")[0] == "Shown by Smith et al. earlier."


def _paper_with(sections):
    return StructuredPaper(source_path="t", title="Test Paper", abstract="Text.", keywords=["deep learning", "CT"],
                           sections=sections)


def test_single_appendix_is_unnumbered():
    paper = map_to_ieee(_paper_with([
        Section(heading="Introduction", number="1", level=1,
                elements=[Element("paragraph", "We study it.")]),
        Section(heading="Extra Results", number="A", level=1, in_appendix=True,
                elements=[Element("paragraph", "More.")]),
    ]))
    assert paper.sections[1].ieee_heading == "Appendix: Extra Results"   # HDG-08


def test_word_file_rules(tmp_path):
    paper = map_to_ieee(_paper_with([
        Section(heading="Acknowledgment", number="", level=1, elements=[Element("paragraph", "Thanks.")]),
        Section(heading="Introduction", number="1", level=1,
                elements=[Element("paragraph", "We study it.")]),
        Section(heading="Training", number="1.1.1", level=3,
                elements=[Element("paragraph", "We trained it.")]),
        Section(heading="Extra Results", number="A", level=1, in_appendix=True,
                elements=[Element("paragraph", "More.")]),
    ]))
    paper.keywords = sorted(paper.keywords, key=str.lower)
    path = str(tmp_path / "rules.docx")
    render_docx(paper, path)
    texts = [p.text for p in docx.Document(path).paragraphs]

    assert "1) Training: We trained it." in texts                              # HDG-03 run-in
    order = [t for t in texts if t in ("Acknowledgment", "References") or t.startswith("Appendix")]
    assert order == ["Appendix: Extra Results", "Acknowledgment", "References"]  # ACK-02
    terms = next(t for t in texts if t.startswith(R.LABELS["index_terms"]))
    assert terms == "Index Terms—CT, deep learning"                            # ABS-07, ABS-08

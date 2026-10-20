"""
Tests for stage 2 (classifier + assembler).

Run with:   pytest -q

Each test builds a small document with a known answer, runs stage 2 on it,
and checks the answer came out. If you change a rule in classifier.py and a
test goes red, the test name tells you which kind of paper you just broke.
"""

from __future__ import annotations

import os
import sys

import docx
import pymupdf
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from src.ingest import load_document                             # noqa: E402
from src.models import BlockType                                 # noqa: E402
from src.structure import structure_document                     # noqa: E402
from src.structure.assembler import join_lines                   # noqa: E402
from src.structure.classifier import split_heading_number        # noqa: E402

SAMPLE_PDF = os.path.join("data", "input", "sample_paper.pdf")


# ---------------------------------------------------------------------------
# Small pure functions
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("text, expected", [
    ("3.1 Dataset", ("3.1", "Dataset", 2)),
    ("1. Background", ("1", "Background", 1)),
    ("IV. RESULTS", ("IV", "RESULTS", 1)),
    ("B. Training Setup", ("B", "Training Setup", 2)),
    ("Results", ("", "Results", 0)),
])
def test_split_heading_number(text, expected):
    assert split_heading_number(text) == expected


def test_join_lines_repairs_hyphenated_words():
    assert join_lines(["deep convolu-", "tional networks"]) == "deep convolutional networks"
    assert join_lines(["state-of-the-", "Art model"]) == "state-of-the- Art model"
    assert join_lines(["one", "two  three"]) == "one two three"


# ---------------------------------------------------------------------------
# The sample paper - single column, APA references
# ---------------------------------------------------------------------------

@pytest.fixture(scope="module")
def sample():
    if not os.path.exists(SAMPLE_PDF):
        pytest.skip("run tools/make_sample_paper.py first")
    doc = load_document(SAMPLE_PDF)
    paper = structure_document(doc)
    return doc, paper


def test_sample_front_matter(sample):
    _, paper = sample
    assert paper.title.startswith("Automated Detection of Pulmonary Nodules")
    assert paper.title.endswith("Convolutional Neural Networks")      # both lines joined
    assert paper.authors == ["K. M. Perera", "S. Fernando", "A. Jayasuriya"]
    assert "University of Colombo" in paper.affiliations[0]
    assert paper.abstract.startswith("Lung cancer remains")
    assert "deep learning" in paper.keywords


def test_sample_running_header_and_page_numbers_are_junk(sample):
    doc, _ = sample
    junk = [b.text for b in doc.blocks if b.block_type is BlockType.JUNK]
    assert junk.count("Preprint - not peer reviewed") == 3       # one per page
    assert "1" in junk and "2" in junk
    # the labels drawn inside Figure 1 are part of the picture, not body text
    assert "Encoder (5 stages)" in junk


def test_sample_sections(sample):
    _, paper = sample
    outline = [(s.number, s.heading, s.level) for s in paper.sections]
    assert outline == [
        ("1", "Background", 1),
        ("2", "Previous Studies", 1),
        ("3", "Our Approach", 1),
        ("3.1", "Dataset", 2),
        ("3.2", "Network Architecture", 2),
        ("4", "Experiments and Findings", 1),
        ("5", "Final Remarks", 1),
    ]


def test_sample_paragraphs_are_rebuilt(sample):
    _, paper = sample
    background = paper.sections[0]
    assert len(background.paragraphs) == 2
    assert background.paragraphs[1].startswith("Computer aided detection")


def test_sample_equation_and_captions_stay_in_their_section(sample):
    _, paper = sample
    architecture = paper.sections[4]
    kinds = [e.kind for e in architecture.elements]
    # In the SOURCE the figure floats above the paragraph that mentions it;
    # Step 6 (placement) moves it after that paragraph - see test_visuals.py
    assert kinds == ["caption", "paragraph", "equation", "paragraph"]
    assert paper.sections[5].elements[1].text.startswith("Table 1")


def test_sample_references_are_merged_into_entries(sample):
    _, paper = sample
    assert len(paper.references) == 4
    assert paper.references[0].startswith("Dou, Q.")
    assert paper.references[3].endswith("209-249.")


# ---------------------------------------------------------------------------
# A two-column, IEEE-flavoured paper built on the fly
# ---------------------------------------------------------------------------

def _two_column_pdf(path: str) -> None:
    """Title across the top, then two columns, the way most papers look.

    The right column deliberately starts higher than some left-column text,
    so a naive top-to-bottom reader would interleave the two.
    """
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    page.insert_textbox(pymupdf.Rect(50, 50, 562, 90),
                        "A Study of Two Column Layouts", fontsize=20,
                        fontname="tibo", align=1)
    page.insert_textbox(pymupdf.Rect(50, 95, 562, 125), "N. Silva and R. Kumar",
                        fontsize=11, fontname="tiro", align=1)

    # The built-in PDF test fonts cannot draw an em dash, so this fixture uses
    # a plain hyphen. The em-dash form "Abstract—" is tested in the Word test.
    left = (
        ("Abstract - We show that column order matters when reading papers.", 9, "tibo"),
        ("Index Terms - layout analysis, reading order", 9, "tibo"),
        ("I. INTRODUCTION", 10, "tiro"),
        ("LEFT ONE. This sentence belongs to the introduction and it lives in "
         "the left column of the page.", 10, "tiro"),
    )
    right = (
        ("II. METHOD", 10, "tiro"),
        ("RIGHT ONE. This sentence belongs to the method section and it lives "
         "in the right column of the page.", 10, "tiro"),
    )

    y = 130
    for text, size, font in left:
        page.insert_textbox(pymupdf.Rect(50, y, 296, y + 60), text, fontsize=size, fontname=font)
        y += 45
    y = 130
    for text, size, font in right:
        page.insert_textbox(pymupdf.Rect(316, y, 562, y + 60), text, fontsize=size, fontname=font)
        y += 45

    # References at the bottom of the right column, numbered IEEE style
    page.insert_textbox(pymupdf.Rect(316, 300, 562, 320), "REFERENCES",
                        fontsize=10, fontname="tiro")
    page.insert_textbox(pymupdf.Rect(316, 320, 562, 360),
                        "[1] A. Author, \"First paper,\" IEEE Trans., 2020.",
                        fontsize=8, fontname="tiro")
    page.insert_textbox(pymupdf.Rect(316, 360, 562, 400),
                        "[2] B. Author, \"Second paper,\" in Proc. CVPR, 2021.",
                        fontsize=8, fontname="tiro")
    doc.save(path)


@pytest.fixture(scope="module")
def two_column(tmp_path_factory):
    path = str(tmp_path_factory.mktemp("pdf") / "two_column.pdf")
    _two_column_pdf(path)
    doc = load_document(path)
    return doc, structure_document(doc)


def test_two_column_reading_order(two_column):
    doc, _ = two_column
    text = " ".join(b.text for b in doc.blocks)
    assert text.index("LEFT ONE") < text.index("II. METHOD") < text.index("RIGHT ONE")


def test_two_column_inline_abstract_and_index_terms(two_column):
    _, paper = two_column
    assert paper.title == "A Study of Two Column Layouts"
    assert paper.authors == ["N. Silva", "R. Kumar"]
    assert paper.abstract.startswith("We show that column order")
    assert paper.keywords == ["layout analysis", "reading order"]


def test_two_column_roman_headings_and_numbered_references(two_column):
    _, paper = two_column
    assert [(s.number, s.heading) for s in paper.sections] == [
        ("I", "INTRODUCTION"), ("II", "METHOD"),
    ]
    assert len(paper.references) == 2
    assert paper.references[1].startswith("[2]")


# ---------------------------------------------------------------------------
# A Word draft using real heading styles
# ---------------------------------------------------------------------------

def test_word_draft(tmp_path):
    path = tmp_path / "draft.docx"
    d = docx.Document()
    d.add_paragraph("Smart Irrigation With Low Cost Sensors", style="Title")
    d.add_paragraph("P. Herath and D. Mendis")
    d.add_paragraph("Abstract—We built a cheap irrigation controller.")
    d.add_heading("Introduction", level=1)
    d.add_paragraph("Water is scarce in the dry zone.")
    d.add_paragraph("Farmers irrigate by guesswork.")
    d.add_heading("Hardware", level=2)
    d.add_paragraph("We used an ESP32 board.")
    d.add_heading("References", level=1)
    d.add_paragraph("[1] X. Y. Zed, \"Sensors,\" 2019.")
    d.add_paragraph("[2] Q. R. Ess, \"Soil,\" 2020.")
    d.save(path)

    paper = structure_document(load_document(str(path)))
    assert paper.title == "Smart Irrigation With Low Cost Sensors"
    assert paper.authors == ["P. Herath", "D. Mendis"]
    assert paper.abstract == "We built a cheap irrigation controller."
    assert [(s.heading, s.level) for s in paper.sections] == [
        ("Introduction", 1), ("Hardware", 2),
    ]
    assert len(paper.sections[0].paragraphs) == 2      # Word paragraphs kept apart
    assert len(paper.references) == 2


def test_word_heading_is_never_taken_as_the_title(tmp_path):
    """A draft with no title: its biggest text is a 'Heading 1', which must
    stay a section heading instead of being promoted to the title."""
    path = tmp_path / "untitled.docx"
    d = docx.Document()
    d.add_heading("Our Approach", level=1)
    d.add_paragraph("We used the LUNA16 collection of 888 thoracic CT scans.")
    d.save(path)

    paper = structure_document(load_document(str(path)))
    assert paper.title == ""
    assert [s.heading for s in paper.sections] == ["Our Approach"]

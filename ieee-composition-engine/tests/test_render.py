"""
Tests for stage 4 (the IEEE Word renderer).

These open the .docx the engine wrote and check the IEEE rules are really in
it - the page size, margins, columns, fonts, headings and captions - rather
than trusting that the code "probably" did it.

Run with:   pytest -q
"""

from __future__ import annotations

import os
import sys

import docx
import pytest
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from src.ingest import load_document                                   # noqa: E402
from src.models import Element, Section, StructuredPaper               # noqa: E402
from src.render import render_docx                                     # noqa: E402
from src.render import ieee_style as S                                 # noqa: E402
from src.structure import map_to_ieee, structure_document              # noqa: E402

SAMPLE_PDF = os.path.join("data", "input", "sample_paper.pdf")


def _columns(section) -> int:
    cols = section._sectPr.find(qn("w:cols"))
    return int(cols.get(qn("w:num"), "1")) if cols is not None else 1


def _find(document, startswith: str):
    for p in document.paragraphs:
        if p.text.startswith(startswith):
            return p
    raise AssertionError(f"no paragraph starting with {startswith!r}")


@pytest.fixture(scope="module")
def rendered(tmp_path_factory):
    if not os.path.exists(SAMPLE_PDF):
        pytest.skip("run tools/make_sample_paper.py first")
    paper = map_to_ieee(structure_document(load_document(SAMPLE_PDF)))
    path = str(tmp_path_factory.mktemp("out") / "paper_IEEE.docx")
    notes = render_docx(paper, path)
    return docx.Document(path), notes


# ---------------------------------------------------------------------------
# Page layout
# ---------------------------------------------------------------------------

def test_page_size_and_margins(rendered):
    document, _ = rendered
    for section in document.sections:
        assert section.page_width == Inches(8.5)
        assert section.page_height == Inches(11)
        assert section.top_margin == Inches(0.75)
        assert section.bottom_margin == Inches(1.0)
        assert section.left_margin == Inches(0.625)
        assert section.right_margin == Inches(0.625)


def test_title_part_is_one_column_body_is_two(rendered):
    document, _ = rendered
    assert _columns(document.sections[0]) == 1
    assert all(_columns(s) == 2 for s in document.sections[1:])


def test_page_width_adds_up():
    total = S.MARGIN_LEFT + S.COLUMN_WIDTH + S.COLUMN_GAP + S.COLUMN_WIDTH + S.MARGIN_RIGHT
    assert total == S.PAGE_WIDTH


# ---------------------------------------------------------------------------
# Text styles
# ---------------------------------------------------------------------------

def test_title(rendered):
    document, _ = rendered
    title = document.paragraphs[0]
    assert title.text.startswith("Automated Detection of Pulmonary Nodules")
    assert title.runs[0].font.size == Pt(24)


def test_every_run_is_times_new_roman(rendered):
    """Every run is Times New Roman - except code listings, which use Courier."""
    document, _ = rendered
    for p in document.paragraphs:
        for run in p.runs:
            assert run.font.name in ("Times New Roman", "Courier New"), p.text[:40]


def test_abstract_and_index_terms(rendered):
    document, _ = rendered
    abstract = _find(document, "Abstract—")
    label, text = abstract.runs[0], abstract.runs[1]
    assert label.font.bold and label.font.italic
    assert text.font.bold and text.font.size == Pt(9)
    assert _find(document, "Index Terms—").text.endswith("medical imaging")


def test_section_headings_are_roman_small_caps(rendered):
    document, _ = rendered
    for label in ("I. Introduction", "II. Related Work", "III. Methodology",
                  "IV. Results", "V. Conclusion"):
        p = _find(document, label)
        assert p.runs[0].font.small_caps, label
        assert p.paragraph_format.keep_with_next, label


def test_subheadings_are_lettered_italics(rendered):
    document, _ = rendered
    p = _find(document, "A. Dataset")
    assert p.runs[0].font.italic
    _find(document, "B. Network Architecture")


def test_body_text(rendered):
    document, _ = rendered
    p = _find(document, "Lung cancer accounted")
    assert p.runs[0].font.size == Pt(10)
    assert p.paragraph_format.first_line_indent == S.BODY_FIRST_LINE_INDENT


def test_text_after_equation_is_not_indented(rendered):
    document, _ = rendered
    p = _find(document, "where alpha was fixed")
    assert not p.paragraph_format.first_line_indent


# ---------------------------------------------------------------------------
# Equations, captions, references
# ---------------------------------------------------------------------------

def test_equation_number_is_tabbed_to_the_right(rendered):
    document, _ = rendered
    p = _find(document, "\tL = alpha")
    assert p.runs[-1].text == "\t(1)"


def test_figure_caption_uses_fig_abbreviation(rendered):
    document, _ = rendered
    p = _find(document, "Fig. 1.")
    # FIG-02: "Fig. 1." then an em space (U+2003), as the IEEE Editorial Style Manual asks
    assert p.text == "Fig. 1.\u2003Overall architecture of the proposed detection network."
    assert p.runs[0].font.size == Pt(8)


def test_table_caption_uses_roman_numeral_and_small_caps(rendered):
    document, _ = rendered
    label = _find(document, "TABLE I")
    assert label.text == "TABLE I"
    caption = document.paragraphs[[p.text for p in document.paragraphs].index("TABLE I") + 1]
    assert caption.runs[0].font.small_caps


def test_references_are_numbered_with_hanging_indent(rendered):
    document, _ = rendered
    refs = [p for p in document.paragraphs if p.text.startswith("[")
            and p.text[1:2].isdigit()]
    assert [r.text[:3] for r in refs] == ["[1]", "[2]", "[3]", "[4]"]
    assert refs[0].paragraph_format.first_line_indent < 0
    assert refs[0].runs[0].font.size == Pt(8)


def test_renderer_reports_reference_style_is_not_yet_converted(rendered):
    _, notes = rendered
    assert any("Step 5" in n for n in notes)


# ---------------------------------------------------------------------------
# A paper with missing parts must still render, with honest placeholders
# ---------------------------------------------------------------------------

def test_missing_parts_get_placeholders_not_invented_text(tmp_path):
    paper = StructuredPaper(source_path="x", title="A Short Draft")
    s = Section(heading="Introduction", number="1", level=1)
    s.elements.append(Element(kind="paragraph", text="Only one paragraph here."))
    paper.sections.append(s)
    map_to_ieee(paper)

    path = tmp_path / "draft.docx"
    notes = render_docx(paper, str(path))
    text = "\n".join(p.text for p in docx.Document(str(path)).paragraphs)

    assert "[Abstract missing" in text
    assert "[Index terms missing" in text
    assert "No references found" in text
    assert any("Abstract missing" in n for n in notes)


def test_already_numbered_references_are_not_double_numbered(tmp_path):
    paper = StructuredPaper(source_path="x", title="T", abstract="A.", keywords=["k"],
                            references=["[1] A. Author, \"Paper,\" 2020.",
                                        "2. B. Author, \"Other,\" 2021."])
    path = tmp_path / "refs.docx"
    render_docx(paper, str(path))
    texts = [p.text for p in docx.Document(str(path)).paragraphs if p.text.startswith("[")]
    assert texts == ["[1]\tA. Author, \"Paper,\" 2020.", "[2]\tB. Author, \"Other,\" 2021."]


def test_control_characters_never_crash_the_renderer(tmp_path):
    paper = StructuredPaper(source_path="x", title="Bad\x01 Title", abstract="A\x0b.",
                            keywords=["k"])
    s = Section(heading="Introduction", number="1", level=1)
    s.elements.append(Element(kind="paragraph", text="Text with \x00 a null and \x1f more."))
    paper.sections.append(s)
    map_to_ieee(paper)
    path = tmp_path / "bad.docx"
    render_docx(paper, str(path))                       # must not raise
    text = "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
    assert "Text with  a null and  more." in text


def test_code_listing_is_kept_in_typewriter_font(tmp_path):
    paper = StructuredPaper(source_path="x", title="T", abstract="A.", keywords=["k"])
    s = Section(heading="Appendix", number="A", level=1, in_appendix=True)
    s.elements.append(Element(kind="code", text='{"action": "look"}\nprint(answer)'))
    paper.sections.append(s)
    map_to_ieee(paper)
    path = tmp_path / "code.docx"
    render_docx(paper, str(path))
    paras = {p.text: p for p in docx.Document(str(path)).paragraphs}
    assert paras["print(answer)"].runs[0].font.name == "Courier New"
    assert '{"action": "look"}' in paras

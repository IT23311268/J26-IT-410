"""
Tests for the rubric-based IEEE evaluation (src/rubric/).

Run with:   pytest -q

The sample paper goes through the whole engine once; the tests then check
that every rule has a rubric line, that the receipts say EXACTLY where each
rule was applied (before -> after), that the score responds to ticks, and
that the Excel download and the annotated Word file contain the same thing.
"""

from __future__ import annotations

import io
import os
import sys

import docx
import pytest
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from src.compose.llm import NoLLM                      # noqa: E402
from src.pipeline import run_pipeline                  # noqa: E402
from src.rubric import rubric_score                    # noqa: E402
from src.rubric.export import rubric_workbook          # noqa: E402
from src.rules import ieee_rules as R                  # noqa: E402

SAMPLE_PDF = os.path.join(ROOT, "data", "input", "sample_paper.pdf")


@pytest.fixture(scope="module")
def result(tmp_path_factory):
    if not os.path.exists(SAMPLE_PDF):
        pytest.skip("run tools/make_sample_paper.py first")
    out = str(tmp_path_factory.mktemp("rubric"))
    return run_pipeline(SAMPLE_PDF, outdir=out, llm=NoLLM("test"))


def by_id(result):
    return {i.id: i for i in result.rubric}


def test_every_rulebook_rule_has_a_rubric_line(result):
    assert [i.id for i in result.rubric] == [r.id for r in R.RULES]


def test_rules_apply_only_when_the_paper_has_that_part(result):
    items = by_id(result)
    assert items["TAB-02"].applies                    # the sample has a table
    assert not items["HDG-03"].applies                # ... but no level-3 headings
    assert items["HDG-03"].state == "na" and "level-3" in items["HDG-03"].reason
    assert items["FIG-09"].state == "review"          # author's judgement: a person checks it


def test_receipts_say_exactly_where(result):
    items = by_id(result)
    fig = items["FIG-02"].evidence[0]
    assert fig["where"] == "Fig. 1 caption"
    assert fig["before"].startswith("Figure 1") and fig["after"] == "Fig. 1."

    heading = next(e for e in items["HDG-01"].evidence if "Background" in e["before"])
    assert heading["after"] == "I. INTRODUCTION"

    cite = items["CIT-01"].evidence[0]
    assert cite["where"].startswith("Section I (Introduction), paragraph")
    assert "2020" in cite["before"] or "2021" in cite["before"]
    assert cite["after"].startswith("[")

    moved = items["FIG-05"].evidence[0]
    assert "BEFORE its first mention" in moved["before"]
    assert moved["after"] == "right after its first mention"

    margins = items["LAY-02"].evidence[0]
    assert margins["before"] and "0.75" in margins["after"]     # source value -> IEEE value


def test_unfixable_rule_is_reported_not_ticked(result):
    abstract = by_id(result)["ABS-03"]               # 83 words: the engine never pads
    assert abstract.state in ("partly", "not_met") and not abstract.suggested


def test_the_score_follows_the_ticks(result):
    ticked, applicable, pct = rubric_score(result.rubric)
    assert applicable == sum(1 for i in result.rubric if i.applies)
    assert ticked == sum(1 for i in result.rubric if i.applies and i.suggested)
    everything = {i.id: True for i in result.rubric}
    assert rubric_score(result.rubric, everything)[2] == 100.0
    fewer = rubric_score(result.rubric, {"LAY-01": False})[2]
    assert fewer < pct
    # a rule that does not apply never counts, ticked or not
    assert rubric_score(result.rubric, {"HDG-03": True}) == (ticked, applicable, pct)


def test_excel_download(result):
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(rubric_workbook(result.rubric, {}, "Sample", 95.0)))
    rows = list(book["Rubric"].iter_rows(values_only=True))
    assert len(rows) == len(R.RULES) + 1
    fig = next(r for r in rows if r[2] == "FIG-02")
    assert fig[0] == "✔" and fig[1] == "Not reviewed" and "Fig. 1 caption" in fig[8]
    # with the reviewer's answers from the guided review
    book = load_workbook(io.BytesIO(rubric_workbook(result.rubric, {"FIG-02": False}, "Sample",
                                                    95.0, reviewed={"FIG-02": False})))
    fig = next(r for r in book["Rubric"].iter_rows(values_only=True) if r[2] == "FIG-02")
    assert not fig[0] and fig[1] == "No"                      # an empty Excel cell reads as None
    summary = {r[0]: r[1] for r in book["Summary"].iter_rows(values_only=True)}
    assert summary["Engine's automatic compliance score"] == "95.0%"


def test_annotated_word_file_has_the_rules_in_the_margin(result):
    document = docx.Document(result.annotated_path)
    comments = list(document.comments)
    assert len(comments) >= 15
    # find the paragraph each comment is attached to
    attached = {}
    for p in document.element.body.iter(qn("w:p")):
        text = "".join(t.text or "" for t in p.iter(qn("w:t")))
        for start in p.iter(qn("w:commentRangeStart")):
            attached[start.get(qn("w:id"))] = text
    caption = next(c for c in comments if "FIG-02" in c.text)
    assert attached[str(caption.comment_id)].startswith("Fig. 1.")
    heading = next(c for c in comments if "HDG-01" in c.text and "Background" in c.text)
    assert attached[str(heading.comment_id)] == "I. Introduction"
    # the plain manuscript has no comments
    assert not list(docx.Document(result.docx_path).comments)


def test_preview_highlights_the_exact_place(result):
    from src.ui.preview import build_preview
    from src.ui.rubric_tab import highlights_for
    items = by_id(result)
    page = build_preview(result.paper, highlights_for(items["FIG-02"], 0))
    assert '<mark class="hl hl-cur" id="hl-current">Fig. 1.</mark>' in page
    cite = items["CIT-01"]
    page = build_preview(result.paper, highlights_for(cite, 0))
    assert f'id="hl-current">{cite.evidence[0]["after"]}</mark>' in page
    # a whole-page rule: the banner explains it and the page stays at the top
    page = build_preview(result.paper, highlights_for(items["LAY-01"], 0), banner="LAY-01")
    assert '<div class="banner" id="hl-current">LAY-01</div>' in page and "hl-page" in page
    # a moved figure: the figure AND the paragraph that mentions it
    page = build_preview(result.paper, highlights_for(items["FIG-05"], 0))
    assert 'class="hl-rel"' in page and "hl-block-cur" in page

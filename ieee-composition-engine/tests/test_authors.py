"""
Tests for the author block (rules TTL-03, TTL-06).

The problem these tests guard against came from a real paper ("Scaling
Discovery Requires Forgetting"): page 1 has the title, the authors with
superscript affiliation numbers, then a FIGURE with its caption, then the
affiliations, then the abstract. The old front-matter rule treated every
line between the title and the abstract as an author, so the graph labels
and the Figure 1 caption ended up in the author list.

Run with:   pytest -q
"""

from __future__ import annotations

import os
import sys

import docx
import pymupdf
import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from src.ingest import load_document                      # noqa: E402
from src.structure import map_to_ieee, structure_document  # noqa: E402
from src.structure.assembler import parse_affiliations, parse_authors  # noqa: E402

AUTHORS = [("Shangyin Tan", "1"), ("Ben Zhou", "1,3"), ("Lakshya A Agrawal", "1"),
           ("Jialin Zhang", "2"), ("Alex Dimakis", "1,3"), ("Matei Zaharia", "1")]


def _teaser_pdf(path: str) -> None:
    """Title, authors with superscripts, a drawn figure + caption, affiliations, abstract."""
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    page.insert_textbox(pymupdf.Rect(72, 60, 540, 90), "Scaling Discovery Requires Forgetting",
                        fontsize=17, fontname="tibo", align=1)

    # authors on two lines, each name followed by superscript affiliation numbers
    x, y = 80, 112
    for i, (name, marks) in enumerate(AUTHORS):
        if i == 4:
            x, y = 150, 126
        page.insert_text((x, y), name, fontsize=10, fontname="tiro")
        x += pymupdf.get_text_length(name, "tiro", 10)
        page.insert_text((x, y - 4), marks, fontsize=6, fontname="tiro")
        x += pymupdf.get_text_length(marks, "tiro", 6)
        if i not in (3, len(AUTHORS) - 1):
            page.insert_text((x, y), ",", fontsize=10, fontname="tiro")
            x += pymupdf.get_text_length(", ", "tiro", 10)

    # the teaser figure: two small plots with axis titles drawn in a sans font
    for left in (110, 330):
        page.draw_rect(pymupdf.Rect(left, 170, left + 170, 270), color=(0, 0, 0), width=0.6)
        page.draw_polyline([(left + 5, 260), (left + 60, 230), (left + 120, 200),
                            (left + 165, 185)], color=(0.2, 0.3, 0.9), width=1.2)
    page.insert_text((115, 163), "AHC014: score, higher is better", fontsize=9, fontname="helv")
    page.insert_text((335, 163), "VLIW kernel: cycles, lower is better", fontsize=9,
                     fontname="helv")
    page.insert_text((180, 282), "Total money spent ($)", fontsize=9, fontname="helv")

    page.insert_textbox(pymupdf.Rect(90, 292, 522, 330),
                        "Figure 1: Best score so far against total money spent on AtCoder "
                        "Heuristic Contest and a VLIW kernel take-home, for the best replicate "
                        "of each configuration.", fontsize=9, fontname="tiro")

    # affiliations with superscript numbers
    y = 345
    for mark, aff in (("1", "UC Berkeley"), ("2", "Stanford"), ("3", "Bespoke Labs")):
        width = pymupdf.get_text_length(aff, "tiit", 9)
        x = 306 - width / 2
        page.insert_text((x - 4, y - 3), mark, fontsize=6, fontname="tiit")
        page.insert_text((x, y), aff, fontsize=9, fontname="tiit")
        y += 12

    page.insert_textbox(pymupdf.Rect(250, 392, 362, 412), "Abstract", fontsize=11,
                        fontname="tibo", align=1)
    page.insert_textbox(pymupdf.Rect(110, 415, 502, 470),
                        "Long-running research agents plateau. We show that discarding the "
                        "session and restarting from the best candidate unblocks progress.",
                        fontsize=10, fontname="tiro")
    page.insert_textbox(pymupdf.Rect(72, 482, 540, 512), "1 Introduction", fontsize=12,
                        fontname="tibo")
    page.insert_textbox(pymupdf.Rect(72, 510, 540, 560),
                        "Agents that keep their whole history get stuck, as Figure 1 shows. "
                        "Forgetting helps them climb again.", fontsize=10, fontname="tiro")
    doc.save(path)


@pytest.fixture(scope="module")
def teaser(tmp_path_factory):
    path = str(tmp_path_factory.mktemp("authors") / "teaser.pdf")
    _teaser_pdf(path)
    doc = load_document(path)
    return path, map_to_ieee(structure_document(doc))


def test_only_real_authors_are_authors(teaser):
    _, paper = teaser
    assert paper.authors == [name for name, _ in AUTHORS]
    text = " ".join(paper.authors)
    assert "AHC014" not in text and "Figure" not in text and "higher" not in text


def test_affiliations_are_clean_and_linked(teaser):
    _, paper = teaser
    assert paper.affiliations == ["UC Berkeley", "Stanford", "Bespoke Labs"]
    # Ben Zhou is 1,3 -> UC Berkeley and Bespoke Labs
    assert paper.author_affiliations[1] == [0, 2]
    assert paper.author_affiliations[3] == [1]


def test_the_front_page_figure_is_a_real_figure(teaser):
    _, paper = teaser
    captions = [e.text for s in paper.sections for e in s.elements if e.kind == "caption"]
    assert any(c.startswith("Figure 1: Best score") for c in captions)
    # the figure before the first heading joins the Introduction - no empty
    # untitled section, so no second "Introduction"
    assert [s.ieee_heading for s in paper.sections] == ["Introduction"]


@pytest.mark.parametrize("lines, names, marks", [
    (["Shangyin Tan1, Ben Zhou1,3, Lakshya A Agrawal1,"], ["Shangyin Tan", "Ben Zhou",
                                                         "Lakshya A Agrawal"],
     [["1"], ["1", "3"], ["1"]]),
    (["K. M. Perera, S. Fernando and A. Jayasuriya"], ["K. M. Perera", "S. Fernando",
                                                       "A. Jayasuriya"], [[], [], []]),
    (["Haofei Yu1 ∗", "Yixing Chen2 ∗"], ["Haofei Yu", "Yixing Chen"], [["1"], ["2"]]),
    (["Alice Smith1, Bob Jones2, Carol", "White1"], ["Alice Smith", "Bob Jones", "Carol White"],
     [["1"], ["2"], ["1"]]),
])
def test_parse_authors(lines, names, marks):
    got_names, got_marks = parse_authors(lines)
    assert got_names == names and got_marks == marks


def test_parse_affiliations():
    names, by_mark = parse_affiliations(["1UC Berkeley 2Stanford", "3Bespoke Labs",
                                         "∗Equal contribution"])
    assert names == ["UC Berkeley", "Stanford", "Bespoke Labs"]
    assert by_mark == {"1": 0, "2": 1, "3": 2}


def test_word_file_lists_each_author_separately(teaser, tmp_path):
    from src.render import render_docx
    _, paper = teaser
    path = str(tmp_path / "authors.docx")
    render_docx(paper, path)
    document = docx.Document(path)
    table = document.tables[0]                                  # TTL-06: the author block
    cells = [c for row in table.rows for c in row.cells if c.text.strip()]
    assert len(cells) == len(AUTHORS)
    ben = next(c for c in cells if c.paragraphs[0].text == "Ben Zhou")
    assert [p.text for p in ben.paragraphs[1:]] == ["UC Berkeley", "Bespoke Labs"]
    assert ben.paragraphs[1].runs[0].font.italic

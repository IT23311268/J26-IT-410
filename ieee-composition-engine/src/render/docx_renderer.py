"""
docx_renderer.py - stage 4: write the IEEE-formatted Word document.

Input : a StructuredPaper that has been through stage 3 (it knows its IEEE
        headings and numbers).
Output: a .docx file laid out the way IEEE requires.

The document is built in two parts, because that is how IEEE pages work:

    PART 1 - one column across the full page
             the title, the authors, the affiliations

    PART 2 - two columns
             abstract, index terms, every section, references

Word calls these parts "sections" (not to be confused with the paper's
sections). Part 2 starts with a "continuous" section break, which means the
column change happens without starting a new page.

All sizes and measurements come from ieee_style.py - this file only decides
WHAT goes on the page, never HOW BIG it is.
"""

from __future__ import annotations

import os
import re
from typing import List, Optional

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Pt, Emu, RGBColor

from ..models import Element, Section, StructuredPaper
from ..structure.ieee_mapper import to_roman
from . import ieee_style as S
from ..rubric.evidence import anchor_of, comments_by_anchor, record

PLACEHOLDER_COLOUR = RGBColor(0x80, 0x80, 0x80)

# Word refuses invisible control characters and crashes if given one.
# The readers already remove them; this is a second safety net.
CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\ufffe\uffff]")


def _add(paragraph, text: str):
    """Add text to a paragraph, minus any characters Word cannot store."""
    return paragraph.add_run(CONTROL_CHARS.sub("", text))

CAPTION_PARTS = re.compile(
    r"^(fig\.?|figure|table|algorithm)\s*([0-9]+|[IVXL]+)[a-z]?\s*[.:\-–—|]?\s*(.*)$", re.I)
EQUATION_NUMBER = re.compile(r"^(.*?)\s*(\(\s*\d+\s*\))\s*$")
ALREADY_NUMBERED_REF = re.compile(r"^\s*(\[\d+\]|\d{1,3}\.)\s*")


# ---------------------------------------------------------------------------
# Low-level helpers
# ---------------------------------------------------------------------------

def _set_font(run, size, bold=False, italic=False, small_caps=False) -> None:
    """Apply IEEE font settings to one run of text.

    Word stores the font name in several places (one per alphabet). If we set
    only the main one, some text silently falls back to Calibri, so we set
    the East-Asian slot too.
    """
    run.font.name = S.FONT
    run.font.size = size
    run.font.bold = bold
    run.font.italic = italic
    run.font.small_caps = small_caps
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(attr), S.FONT)


def _paragraph(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=0,
               first_line=None, keep_with_next=False):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.alignment = align
    fmt.space_before = before
    fmt.space_after = after
    fmt.line_spacing = 1.0
    if first_line is not None:
        fmt.first_line_indent = first_line
    fmt.keep_with_next = keep_with_next
    return p


def _set_columns(section, count: int, gap) -> None:
    """Word has no python-docx setting for columns, so we write the XML tag
    <w:cols w:num="2" w:space="360"/> ourselves. 360 twips = 0.25 inch."""
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), str(int(gap.twips)))


def _page_setup(section) -> None:
    section.page_width = S.PAGE_WIDTH
    section.page_height = S.PAGE_HEIGHT
    section.top_margin = S.MARGIN_TOP
    section.bottom_margin = S.MARGIN_BOTTOM
    section.left_margin = S.MARGIN_LEFT
    section.right_margin = S.MARGIN_RIGHT


def _placeholder(doc, text: str) -> None:
    """A grey, bracketed note marking something a later step will fill in.
    Better an honest gap than invented content."""
    p = _paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=S.CAPTION_SPACE_BEFORE)
    run = _add(p, f"[{text}]")
    _set_font(run, S.CAPTION_SIZE, italic=True)
    run.font.color.rgb = PLACEHOLDER_COLOUR


# ---------------------------------------------------------------------------
# The renderer
# ---------------------------------------------------------------------------

class IEEEDocxRenderer:
    def __init__(self, paper: StructuredPaper, annotate: bool = False):
        self.paper = paper
        self.doc = Document()
        self.notes: List[str] = []       # what the renderer had to work around
        self._run_in = None              # a level-3/4 heading waiting for its paragraph
        # annotate=True writes a copy with every rule receipt as a margin comment;
        # the normal run records the renderer's own receipts instead
        self.annotate = annotate
        self.comments = comments_by_anchor(paper) if annotate else {}
        self._pending_anchor = None      # a run-in heading's comment waits for its paragraph
        self._current_anchor = "document"

    # -- rubric receipts and margin comments ----------------------------------

    def _receipt(self, rules, where: str, before: str = "", after: str = "", note: str = "",
                 anchor: str = "document") -> None:
        if not self.annotate:
            record(self.paper, rules, where, before, after, note=note, anchor=anchor)

    def _comment(self, paragraph, *anchors: str) -> None:
        """Put the receipts for these anchors in the margin next to this paragraph."""
        if not self.annotate or paragraph is None:
            return
        lines = []
        for anchor in anchors:
            lines.extend(self.comments.pop(anchor, []))
        runs = [r for r in paragraph.runs]
        if lines and runs:
            self.doc.add_comment(runs, text="\n".join(lines[:25]) +
                                 (f"\n… and {len(lines) - 25} more" if len(lines) > 25 else ""),
                                 author="IEEE Manuscript Engine", initials="IEEE")

    def _mark(self) -> int:
        """Where the next paragraph will go. Word keeps its section settings
        (w:sectPr) as the LAST item, and new paragraphs are inserted before it."""
        body = self.doc.element.body
        return len(body) - (1 if len(body) and body[-1].tag == qn("w:sectPr") else 0)

    def _first_paragraph_since(self, mark: int, prefer: Optional[str] = None):
        """The first paragraph written since `mark` - or, with `prefer`, the first
        one whose text matches it (a figure's caption, not its picture)."""
        found = None
        for child in self.doc.element.body[mark:]:
            if child.tag != qn("w:p"):
                continue
            paragraph = Paragraph(child, self.doc._body)
            if prefer is None:
                return paragraph
            if re.match(prefer, paragraph.text):
                return paragraph
            found = found or paragraph
        return found

    # -- public -------------------------------------------------------------

    def render(self, output_path: str) -> List[str]:
        self._setup_default_style()
        _page_setup(self.doc.sections[0])
        _set_columns(self.doc.sections[0], 1, S.COLUMN_GAP)

        self._title_block()

        two_col = self.doc.add_section(WD_SECTION.CONTINUOUS)
        _page_setup(two_col)
        _set_columns(two_col, S.COLUMN_COUNT, S.COLUMN_GAP)

        self._abstract_and_index_terms()
        ordered = self._ordered_sections()
        if [s.uid or id(s) for s in ordered] != [s.uid or id(s) for s in self.paper.sections]:
            back = [s for s in ordered if s.level == 1 and s.ieee_slot in ("appendix",
                                                                         "acknowledgment")]
            author_order = [s.ieee_heading for s in self.paper.sections
                            if s.level == 1 and s.ieee_slot in ("appendix", "acknowledgment")]
            self._receipt("ACK-02", "Back matter", " → ".join(author_order + ["References"]),
                          " → ".join([s.ieee_heading for s in back] + ["References"]),
                          anchor=anchor_of(back[0]) if back else "document")
        for section in ordered:
            self._section(section)
        self._references()
        self._document_receipts()

        # IEEE wants the two columns on the last page to end at the same
        # height. Word does this automatically when the two-column part is
        # closed by one more continuous section break.
        closing = self.doc.add_section(WD_SECTION.CONTINUOUS)
        _page_setup(closing)
        _set_columns(closing, S.COLUMN_COUNT, S.COLUMN_GAP)

        self.doc.save(output_path)
        return self.notes

    def _document_receipts(self) -> None:
        """Rules that shape the whole layout - recorded once, with their values, each
        pointing at the part of the manuscript it shapes."""
        p = S.PAPER
        self._receipt("LAY-01", "Whole document", "", f"{p.name} page")
        self._receipt("LAY-02", "Whole document", "", f"margins {p.margin_top_in:.3g} / "
                      f"{p.margin_bottom_in:.3g} / {p.margin_side_in:.3g} in (top/bottom/sides)")
        self._receipt(["LAY-03", "LAY-04"], "Whole document", "",
                      f"{S.COLUMN_COUNT} columns, {p.column_gap_in:.3g} in gap")
        self._receipt("LAY-05", "Title and authors", "", "full page width, above the columns",
                      anchor="front:title")
        self._receipt("LAY-06", "Whole document", "", "justified columns", anchor="all:paragraph")
        self._receipt("LAY-07", "Last page", "", "columns balanced by a closing section break",
                      anchor="references")
        self._receipt("LAY-08", "Whole document", "", "no page numbers, headers or footers")
        self._receipt("TYP-01", "Whole document", "", S.FONT)
        for rule, name, size, anchor in (
                ("TYP-02", "Title", S.TITLE_SIZE, "front:title"),
                ("TYP-03", "Author names", S.AUTHOR_SIZE, "front:authors"),
                ("TYP-04", "Affiliations", S.AFFILIATION_SIZE, "front:authors"),
                ("TYP-05", "Main text", S.BODY_SIZE, "all:paragraph"),
                ("TYP-06", "Abstract and index terms", S.ABSTRACT_SIZE, "front:abstract"),
                ("TYP-07", "Section headings", S.HEADING_SIZE, "all:heading"),
                ("TYP-08", "Captions", S.CAPTION_SIZE, "all:caption"),
                ("TYP-09", "References", S.REFERENCE_SIZE, "references")):
            self._receipt(rule, name, "", f"{size.pt:g} pt", anchor=anchor)
        self._receipt("TYP-10", "Whole document", "", "italics for emphasis, no underlining")
        self._receipt("TYP-11", "Body paragraphs", "",
                      f"first-line indent {p.first_line_indent_in:.2g} in", anchor="all:paragraph")

    # -- document defaults --------------------------------------------------

    def _setup_default_style(self) -> None:
        normal = self.doc.styles["Normal"]
        normal.font.name = S.FONT
        normal.font.size = S.BODY_SIZE
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), S.FONT)
        normal.paragraph_format.space_after = 0
        normal.paragraph_format.line_spacing = 1.0
        # python-docx starts every document with one empty paragraph; reuse
        # nothing from it - we remove it so the title is the first line.
        body = self.doc.element.body
        for p in list(body.iterchildren(qn("w:p"))):
            body.remove(p)

    # -- part 1: title, authors ---------------------------------------------

    def _title_block(self) -> None:
        p = _paragraph(self.doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=S.TITLE_SPACE_AFTER)
        title = self.paper.title or "Untitled Manuscript"
        if not self.paper.title:
            self.notes.append("No title found - used 'Untitled Manuscript'.")
        _set_font(_add(p, title), S.TITLE_SIZE)
        self._comment(p, "front:title", "document")

        if self.paper.authors:
            shown = self._author_block()
        else:
            shown = set()
            self.notes.append("No authors found.")

        # affiliations not printed under any author (none linked, or e-mail lines)
        rest = [a for i, a in enumerate(self.paper.affiliations) if i not in shown]
        for i, affiliation in enumerate(rest):
            last = i == len(rest) - 1
            p = _paragraph(self.doc, align=WD_ALIGN_PARAGRAPH.CENTER,
                           after=S.AUTHOR_BLOCK_SPACE_AFTER if last else 0)
            _set_font(_add(p, affiliation), S.AFFILIATION_SIZE, italic=True)
        if self.paper.authors and not rest:
            spacer = _paragraph(self.doc, after=0)
            spacer.paragraph_format.space_after = S.AUTHOR_BLOCK_SPACE_AFTER

        if not self.paper.affiliations:
            self.doc.paragraphs[-1].paragraph_format.space_after = S.AUTHOR_BLOCK_SPACE_AFTER

    def _author_block(self) -> set:
        """TTL-06: the IEEE author block - every author in a column of their own,
        their affiliation(s) in italics underneath, up to three authors per row.
        Returns the indexes of the affiliations it printed."""
        authors = self.paper.authors
        links = self.paper.author_affiliations
        if len(links) != len(authors):
            links = [[] for _ in authors]
        per_row = min(S.AUTHORS_PER_ROW, len(authors))
        rows = (len(authors) + per_row - 1) // per_row
        table = self.doc.add_table(rows=rows, cols=per_row)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        width = S.PAGE_WIDTH - S.MARGIN_LEFT - S.MARGIN_RIGHT
        shown: set = set()
        for i, name in enumerate(authors):
            cell = table.cell(i // per_row, i % per_row)
            cell.width = width // per_row
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4) if i >= per_row else 0
            _set_font(_add(p, name), S.AUTHOR_SIZE)
            for index in links[i]:
                if 0 <= index < len(self.paper.affiliations):
                    q = cell.add_paragraph()
                    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    _set_font(_add(q, self.paper.affiliations[index]), S.AFFILIATION_SIZE,
                              italic=True)
                    shown.add(index)
        # an uneven last row: centre it by merging the empty cells away is not
        # possible in Word, so the empty cells simply stay blank
        linked = sum(1 for links_ in links if links_)
        self._receipt("TTL-06", "Authors", ", ".join(authors),
                      f"{len(authors)} authors in separate columns, {linked} with their "
                      f"affiliation underneath", anchor="front:authors")
        self._comment(table.cell(0, 0).paragraphs[0], "front:authors")
        return shown

    # -- abstract and index terms -------------------------------------------

    def _abstract_and_index_terms(self) -> None:
        # IEEE: 9pt bold, the label in bold italic, justified.
        p = _paragraph(self.doc, after=S.HEADING2_SPACE_AFTER)
        _set_font(_add(p, S.ABSTRACT_LABEL), S.ABSTRACT_SIZE, bold=True, italic=True)
        if self.paper.abstract:
            _set_font(_add(p, self.paper.abstract), S.ABSTRACT_SIZE, bold=True)
            self._receipt(["ABS-01", "ABS-02"], "Abstract", "", f"“{S.ABSTRACT_LABEL}” run-in "
                          "label in bold italic, one paragraph", anchor="front:abstract")
            self._comment(p, "front:abstract")
        else:
            run = _add(p, "[Abstract missing - the author should add a 150-250 word abstract.]")
            _set_font(run, S.ABSTRACT_SIZE, bold=True)
            run.font.color.rgb = PLACEHOLDER_COLOUR
            self.notes.append("Abstract missing - placeholder inserted.")

        p = _paragraph(self.doc, after=S.HEADING1_SPACE_BEFORE)
        _set_font(_add(p, S.INDEX_TERMS_LABEL), S.ABSTRACT_SIZE, bold=True, italic=True)
        if self.paper.keywords:
            # ABS-07/ABS-08: alphabetical, lower case except acronyms and proper
            # nouns - and only the first word of the list is capitalised.
            terms = ", ".join(self.paper.keywords)
            terms = terms[:1].upper() + terms[1:]
            _set_font(_add(p, terms), S.ABSTRACT_SIZE, bold=True)
            self._receipt(["ABS-06", "ABS-08"], "Index Terms", ", ".join(self.paper.keywords),
                          f"{S.INDEX_TERMS_LABEL}{terms}", note="first word capitalised",
                          anchor="front:index_terms")
            self._comment(p, "front:index_terms")
        else:
            run = _add(p, "[Index terms missing - the author should add 3-5 keywords.]")
            _set_font(run, S.ABSTRACT_SIZE, bold=True)
            run.font.color.rgb = PLACEHOLDER_COLOUR
            self.notes.append("Index terms missing - placeholder inserted.")

    # -- sections -----------------------------------------------------------

    def _heading(self, section: Section) -> None:
        if section.level == 1:
            # "I. Introduction" in small caps, centred.
            p = _paragraph(self.doc, align=WD_ALIGN_PARAGRAPH.CENTER,
                           before=S.HEADING1_SPACE_BEFORE, after=S.HEADING1_SPACE_AFTER,
                           keep_with_next=True)
            label = f"{section.ieee_number}. {section.ieee_heading}" if section.ieee_number \
                else section.ieee_heading
            _set_font(_add(p, label), S.HEADING_SIZE, small_caps=True)
        elif section.level == 2:
            # "A. Dataset" in italics, left-aligned.
            p = _paragraph(self.doc, align=WD_ALIGN_PARAGRAPH.LEFT,
                           before=S.HEADING2_SPACE_BEFORE, after=S.HEADING2_SPACE_AFTER,
                           keep_with_next=True)
            _set_font(_add(p, f"{section.ieee_number}. {section.ieee_heading}"),
                      S.HEADING_SIZE, italic=True)
        else:
            # HDG-03 / HDG-04: "1) Name:" (level 3) or "a) Name:" (level 4) in
            # italics, RUN INTO the first paragraph of the section. Level 4 is
            # indented further. If the section starts with no paragraph, the
            # heading stands on its own line.
            label = f"{section.ieee_number} {section.ieee_heading}:"
            first = next((e for e in section.elements), None)
            if first is not None and first.kind == "paragraph":
                self._run_in = (label, section.level)
                return
            indent = S.BODY_FIRST_LINE_INDENT * (2 if section.level >= 4 else 1)
            p = _paragraph(self.doc, align=WD_ALIGN_PARAGRAPH.LEFT,
                           first_line=indent, keep_with_next=True)
            _set_font(_add(p, label), S.HEADING_SIZE, italic=True)

    def _ordered_sections(self) -> List[Section]:
        return ordered_sections(self.paper)

    def _section(self, section: Section) -> None:
        mark = self._mark()
        self._heading(section)
        if self.annotate:
            heading = self._first_paragraph_since(mark)
            if heading is not None:
                self._comment(heading, anchor_of(section))
            else:                                    # run-in: the comment waits for the text
                self._pending_anchor = anchor_of(section)
        previous = ""
        for element in section.elements:
            mark = self._mark()
            self._element(section, element, previous)
            if self.annotate:
                prefer = (rf"^({re.escape(S.FIGURE_LABEL)}|{S.TABLE_LABEL}|{S.ALGORITHM_LABEL})"
                          if element.kind == "caption" else None)
                first = self._first_paragraph_since(mark, prefer)
                anchors = [anchor_of(element)]
                if self._pending_anchor:
                    anchors.insert(0, self._pending_anchor)
                    self._pending_anchor = None
                self._comment(first, *anchors)
            previous = element.kind

    def _element(self, section: Section, element: Element, previous: str) -> None:
        self._current_anchor = anchor_of(element)
        if element.kind == "paragraph":
            # "where alpha was fixed..." after an equation continues the
            # same sentence, so IEEE does not indent it.
            continues = previous == "equation" and element.text[:1].islower()
            self._body_paragraph(element.text, indent=not continues)
        elif element.kind == "equation":
            self._equation(element.text)
        elif element.kind == "code":
            self._code(element.text)
        elif element.kind == "math_region":
            # A display equation the PDF stored as scattered symbols. It
            # cannot be rebuilt from text, so Step 6 copied it from the
            # source PDF as a picture. Never print the scraps.
            if element.image and os.path.exists(element.image):
                p = _paragraph(self.doc, align=WD_ALIGN_PARAGRAPH.CENTER,
                               before=S.CAPTION_SPACE_BEFORE, after=S.CAPTION_SPACE_BEFORE)
                self._picture(p, element.image)
            else:
                number = re.search(r"\((\d{1,3})\)", element.text)
                label = f"Equation ({number.group(1)})" if number else "Display equation"
                _placeholder(self.doc, f"{label} - could not be recovered from the source")
        elif element.kind == "caption":
            self._caption(element, section)

    def _code(self, text: str) -> None:
        """A code or prompt listing: typewriter font, small, line for line."""
        lines = text.split("\n")
        for i, line in enumerate(lines):
            p = _paragraph(self.doc, align=WD_ALIGN_PARAGRAPH.LEFT,
                           before=S.CAPTION_SPACE_BEFORE if i == 0 else 0,
                           after=S.CAPTION_SPACE_BEFORE if i == len(lines) - 1 else 0)
            run = _add(p, line)
            _set_font(run, S.CODE_SIZE)
            run.font.name = S.CODE_FONT
            r_fonts = run._element.get_or_add_rPr().find(qn("w:rFonts"))
            for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                r_fonts.set(qn(attr), S.CODE_FONT)

    def _body_paragraph(self, text: str, indent: bool = True) -> None:
        run_in, self._run_in = self._run_in, None
        if run_in:                                   # HDG-03 / HDG-04 run-in heading
            label, level = run_in
            p = _paragraph(self.doc, first_line=S.BODY_FIRST_LINE_INDENT * (2 if level >= 4 else 1))
            _set_font(_add(p, label + " "), S.HEADING_SIZE, italic=True)
            _set_font(_add(p, text), S.BODY_SIZE)
            return
        p = _paragraph(self.doc, first_line=S.BODY_FIRST_LINE_INDENT if indent else 0)
        _set_font(_add(p, text), S.BODY_SIZE)

    def _equation(self, text: str) -> None:
        """Equation centred in the column, its number pushed to the right edge.

        Done with two tab stops: a centre tab in the middle of the column and
        a right tab at its edge. This is how the IEEE Word template does it.
        """
        match = EQUATION_NUMBER.match(text)
        body, number = (match.group(1).strip(), match.group(2).replace(" ", "")) \
            if match else (text.strip(), "")

        p = _paragraph(self.doc, align=WD_ALIGN_PARAGRAPH.LEFT,
                       before=S.CAPTION_SPACE_BEFORE, after=S.CAPTION_SPACE_BEFORE)
        stops = p.paragraph_format.tab_stops
        stops.add_tab_stop(S.COLUMN_WIDTH // 2, WD_TAB_ALIGNMENT.CENTER)
        stops.add_tab_stop(S.COLUMN_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
        _set_font(_add(p, "\t" + body), S.BODY_SIZE, italic=True)
        if number:
            _set_font(_add(p, "\t" + number), S.BODY_SIZE)
            self._receipt("EQN-01", f"Equation {number}", text[:60], f"{number} flush right",
                          anchor=self._current_anchor)

    def _caption(self, element: Element, section: Optional[Section] = None) -> None:
        """Figures:  'Fig. 1. Caption text.'       - BELOW the figure
           Tables:   'TABLE I' / 'CAPTION TEXT'    - ABOVE the table, small caps
        """
        match = CAPTION_PARTS.match(element.text)
        if not match:
            p = _paragraph(self.doc, align=WD_ALIGN_PARAGRAPH.CENTER,
                           before=S.CAPTION_SPACE_BEFORE, after=S.CAPTION_SPACE_AFTER)
            _set_font(_add(p, element.text), S.CAPTION_SIZE)
            return

        kind, number, rest = match.group(1).lower(), match.group(2), match.group(3).strip()
        rest = rest.rstrip(".")
        has_image = bool(element.image) and os.path.exists(element.image)

        if kind.startswith("fig"):
            # IEEE: the figure first, its caption BELOW it
            if has_image:
                p = _paragraph(self.doc, align=WD_ALIGN_PARAGRAPH.CENTER,
                               before=S.CAPTION_SPACE_BEFORE, keep_with_next=True)
                self._picture(p, element.image)
            else:
                _placeholder(self.doc, f"Figure {number} image - not found in the source")
            p = _paragraph(self.doc, align=WD_ALIGN_PARAGRAPH.CENTER,
                           before=S.CAPTION_SPACE_BEFORE, after=S.CAPTION_SPACE_AFTER)
            # FIG-02: "Fig. 1." then an em space, then the caption text
            _set_font(_add(p, f"{S.FIGURE_LABEL} {number}.{S.CAPTION_GAP}"), S.CAPTION_SIZE)
            _set_font(_add(p, f"{rest}."), S.CAPTION_SIZE)
            self._receipt(["FIG-01", "FIG-02"], f"{S.FIGURE_LABEL} {number} caption",
                          element.text[:len(match.group(0)) - len(match.group(3))].strip() or
                          element.text[:12], f"{S.FIGURE_LABEL} {number}.",
                          note="caption below the figure", anchor=anchor_of(element))
        elif kind == "algorithm":
            # caption line, then the algorithm box exactly as in the source
            p = _paragraph(self.doc, align=WD_ALIGN_PARAGRAPH.LEFT,
                           before=S.CAPTION_SPACE_BEFORE, keep_with_next=True)
            _set_font(_add(p, f"{S.ALGORITHM_LABEL} {number}. "), S.CAPTION_SIZE, bold=True)
            _set_font(_add(p, rest), S.CAPTION_SIZE)
            if has_image:
                p = _paragraph(self.doc, align=WD_ALIGN_PARAGRAPH.CENTER,
                               after=S.CAPTION_SPACE_AFTER)
                self._picture(p, element.image)
            else:
                _placeholder(self.doc, f"Algorithm {number} - not found in the source")
        else:
            # IEEE: "TABLE II" and its title in small caps ABOVE the table
            roman = number if not number.isdigit() else to_roman(int(number))
            p = _paragraph(self.doc, align=WD_ALIGN_PARAGRAPH.CENTER,
                           before=S.CAPTION_SPACE_BEFORE, keep_with_next=True)
            _set_font(_add(p, f"{S.TABLE_LABEL} {roman}"), S.CAPTION_SIZE)
            p = _paragraph(self.doc, align=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True)
            _set_font(_add(p, rest), S.CAPTION_SIZE, small_caps=True)
            self._receipt(["TAB-01", "TAB-02", "TAB-03"], f"Table {roman} caption",
                          element.text[:60], f"{S.TABLE_LABEL} {roman} / {rest[:50]}",
                          note="title above the table, small capitals, no full stop",
                          anchor=anchor_of(element))
            if has_image:
                p = _paragraph(self.doc, align=WD_ALIGN_PARAGRAPH.CENTER,
                               before=S.CAPTION_SPACE_BEFORE)
                self._picture(p, element.image)
            else:
                _placeholder(self.doc, f"Table {roman} content - not found in the source")
            self.doc.paragraphs[-1].paragraph_format.space_after = S.CAPTION_SPACE_AFTER

    def _picture(self, paragraph, path: str) -> None:
        """Insert an image at its natural size, but never wider than a column.

        Images cropped from the PDF are rendered at 200 dpi, so their natural
        width in inches is pixels / 200. A small equation stays small; a big
        figure is shrunk to fit the 3.5-inch column.
        """
        try:
            import pymupdf
            natural = Emu(int(pymupdf.Pixmap(path).width / 200 * 914400))
            width = min(natural, S.COLUMN_WIDTH)
        except Exception:
            width = S.COLUMN_WIDTH
        paragraph.add_run().add_picture(path, width=width)

    # -- references ---------------------------------------------------------

    def _references(self) -> None:
        p = _paragraph(self.doc, align=WD_ALIGN_PARAGRAPH.CENTER,
                       before=S.HEADING1_SPACE_BEFORE, after=S.HEADING1_SPACE_AFTER,
                       keep_with_next=True)
        _set_font(_add(p, S.REFERENCES_HEADING), S.HEADING_SIZE, small_caps=True)
        if self.paper.references:
            self._receipt("REF-01", "Reference list", "",
                          f"{len(self.paper.references)} references numbered [1], [2]… with a "
                          "hanging indent", anchor="references")
        self._comment(p, "references")

        if not self.paper.references:
            _placeholder(self.doc, "No references found in the source document")
            self.notes.append("No references found.")
            return

        records = self.paper.reference_records
        for i, ref in enumerate(self.paper.references, start=1):
            # Hanging indent: "[1]" sits in the margin, the wrapped lines line up
            # with the start of the text - the IEEE reference-list look.
            p = _paragraph(self.doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            fmt = p.paragraph_format
            fmt.left_indent = S.REFERENCE_HANGING_INDENT
            fmt.first_line_indent = -S.REFERENCE_HANGING_INDENT
            fmt.tab_stops.add_tab_stop(S.REFERENCE_HANGING_INDENT)
            _set_font(_add(p, f"[{i}]\t"), S.REFERENCE_SIZE)

            if i <= len(records) and records[i - 1].get("segments"):
                # Step 5 has split the reference into pieces; the journal or
                # conference name is printed in italics, as IEEE requires.
                for text, italic in records[i - 1]["segments"]:
                    _set_font(_add(p, text), S.REFERENCE_SIZE, italic=italic)
            else:
                _set_font(_add(p, ALREADY_NUMBERED_REF.sub("", ref, count=1)), S.REFERENCE_SIZE)
            self._comment(p, f"reference:{i}")

        if not records:
            self.notes.append(
                "References are numbered [1], [2]... but still in the author's original "
                "style (Step 5 was skipped)."
            )


def ordered_sections(paper: StructuredPaper) -> List[Section]:
    """ACK-02: the main sections first, then appendices, then the Acknowledgment
    (References follow). Subsections stay with their parent. Used by the Word
    file and by the web page's manuscript preview, so both show the same order."""
    groups: List[List[Section]] = []
    for s in paper.sections:
        if s.level == 1 or not groups:
            groups.append([s])
        else:
            groups[-1].append(s)
    rank = {"appendix": 1, "acknowledgment": 2}
    groups.sort(key=lambda g: rank.get(g[0].ieee_slot, 0))     # stable: keeps author order
    return [s for g in groups for s in g]


def render_docx(paper: StructuredPaper, output_path: str, annotate: bool = False) -> List[str]:
    """Write the paper as an IEEE-formatted .docx. Returns notes about any gaps.

    annotate=True writes the same manuscript with every IEEE rule receipt as a
    margin comment at the exact place it was applied (for the rubric)."""
    return IEEEDocxRenderer(paper, annotate=annotate).render(output_path)

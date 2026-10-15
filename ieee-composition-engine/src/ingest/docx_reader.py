"""
docx_reader.py - stage 1b: turn a Word document into TextBlock objects.

Word is friendlier than PDF: it already knows what a paragraph is, and it
often knows what a heading is (the "Heading 1" style). But student drafts are
usually typed with manual bold instead of real styles, so we collect the same
clues as the PDF reader - size, bold, italic - and let stage 2 decide.

Word has no pages in its XML (pagination happens at print time), so page is
always 1 and the x/y coordinates are 0. Stage 2 must therefore not depend on
coordinates being meaningful. It uses them only as a tie-breaker.
"""

from __future__ import annotations

import docx

from ..models import ParsedDocument, TextBlock

DEFAULT_FONT_SIZE = 11.0   # Word's default when a paragraph inherits its size


def _paragraph_font_size(paragraph) -> float:
    """Find a paragraph's font size in points.

    Word stores size on the run if the author set it directly, on the
    paragraph's style if they used a style, and nowhere at all if they left
    the default. We check each place in turn.
    """
    for run in paragraph.runs:
        if run.font.size is not None:
            return round(run.font.size.pt, 1)

    style = paragraph.style
    while style is not None:
        if style.font.size is not None:
            return round(style.font.size.pt, 1)
        style = style.base_style

    return DEFAULT_FONT_SIZE


def _paragraph_is_bold(paragraph) -> bool:
    """Bold if the majority of the characters in the paragraph are bold."""
    bold_chars = sum(len(r.text) for r in paragraph.runs if r.font.bold)
    total_chars = sum(len(r.text) for r in paragraph.runs)
    if total_chars == 0:
        return False
    return bold_chars / total_chars > 0.5


def _paragraph_is_italic(paragraph) -> bool:
    italic_chars = sum(len(r.text) for r in paragraph.runs if r.font.italic)
    total_chars = sum(len(r.text) for r in paragraph.runs)
    if total_chars == 0:
        return False
    return italic_chars / total_chars > 0.5


def read_docx(path: str) -> ParsedDocument:
    """Open a .docx file and return its paragraphs as TextBlock objects."""
    document = docx.Document(path)
    parsed = ParsedDocument(
        source_path=path,
        source_format="docx",
        page_count=1,
    )

    for paragraph in document.paragraphs:
        text = " ".join(paragraph.text.split())
        if not text:
            continue

        style_name = paragraph.style.name if paragraph.style is not None else ""
        size = _paragraph_font_size(paragraph)

        # A paragraph styled as a real Word heading is a strong signal. We
        # nudge its recorded size up so the classifier's "bigger than body
        # text" rule catches it even when the author never changed the size.
        if style_name.lower().startswith("heading") or style_name.lower() == "title":
            size = max(size, DEFAULT_FONT_SIZE + 3)

        parsed.blocks.append(
            TextBlock(
                text=text,
                page=1,
                font_size=size,
                font_name=style_name,
                bold=_paragraph_is_bold(paragraph),
                italic=_paragraph_is_italic(paragraph),
                x0=0.0,
                y0=0.0,
                x1=0.0,
                y1=0.0,
                style_name=style_name,
            )
        )

    parsed.measure_body_font()
    return parsed

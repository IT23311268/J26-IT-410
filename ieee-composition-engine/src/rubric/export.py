"""
export.py - the filled IEEE rubric as an Excel file (rubric download).

Two sheets:
    Summary   the paper, the engine's compliance score, the reviewer's rubric score
    Rubric    one row per rule: ID, rule, what the engine did, where (the
              exact places, before -> after), the reviewer's tick, the IEEE source
"""

from __future__ import annotations

import io
from typing import Dict, List, Optional

from .rubric import RubricItem, rubric_score

STATUS = {"enforced": "Applied automatically", "checked": "Checked and reported",
          "author": "Author's judgement"}


def _places(item: RubricItem, limit: int = 12) -> str:
    lines = []
    for e in item.evidence[:limit]:
        change = f"“{e['before']}” → “{e['after']}”" if e["before"] else e["after"]
        lines.append(f"{e['where']}: {change}".strip(": "))
    if len(item.evidence) > limit:
        lines.append(f"… and {len(item.evidence) - limit} more")
    return "\n".join(lines)


def rubric_workbook(items: List[RubricItem], ticks: Optional[Dict[str, bool]] = None,
                    paper_title: str = "", engine_score: Optional[float] = None,
                    reviewed: Optional[Dict[str, bool]] = None) -> bytes:
    """`reviewed` holds the reviewer's Yes (True) / No (False) answers from the
    guided review; rules missing from it are "Not reviewed"."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    ticks = ticks or {}
    ticked, applicable, pct = rubric_score(items, ticks)
    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    rows = [
        ("IEEE rubric evaluation", ""),
        ("Paper", paper_title),
        ("Rules that apply to this paper", applicable),
        ("Rules ticked by the reviewer", ticked),
        ("Rubric score (ticked / applicable)", f"{pct}%" if pct is not None else "n/a"),
        ("Engine's automatic compliance score",
         f"{engine_score}%" if engine_score is not None else "n/a"),
        ("Rules that do not apply", sum(1 for i in items if not i.applies)),
    ]
    for r in rows:
        summary.append(r)
    summary["A1"].font = Font(bold=True, size=14)
    summary.column_dimensions["A"].width = 40
    summary.column_dimensions["B"].width = 60

    sheet = wb.create_sheet("Rubric")
    header = ["Ticked", "Reviewer answer", "Rule ID", "Group", "Rule", "Applies", "Engine result",
              "Reason",
              "Where it was applied (before → after)", "Rulebook status", "IEEE source"]
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
    for i in items:
        answer = "" if not i.applies else ("Not reviewed" if reviewed is None or
                                            i.id not in reviewed else
                                            "Yes" if reviewed[i.id] else "No")
        sheet.append(["✔" if (i.applies and ticks.get(i.id, i.suggested)) else "", answer,
                      i.id, i.group, i.rule, "yes" if i.applies else "no", i.state_text,
                      i.reason, _places(i), STATUS.get(i.status, i.status), i.source])
    widths = [8, 14, 9, 18, 50, 8, 20, 30, 70, 20, 40]
    for col, width in zip("ABCDEFGHIJK", widths):
        sheet.column_dimensions[col].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "D2"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
